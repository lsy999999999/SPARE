"""Export the selected PDF to editable PowerPoint shapes and text.

Run with simulator/bin/python. Slide 2 preserves the original as a reference.
"""
from pathlib import Path
import json
import re

import pymupdf
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR
from pptx.oxml.xmlchemy import OxmlElement
from pptx.util import Pt

ROOT = Path(__file__).resolve().parent
SCALE = 2.0
# DrawingML's text baseline includes the font's ascent. Calibrated against
# LibreOffice rendering below; all source coordinates are PDF points.
BASELINE_ASCENT = 1.0


def emu(points):
    return round(points * SCALE * 12700)


def element(tag, **attrs):
    node = OxmlElement('a:' + tag)
    for key, value in attrs.items():
        node.set(key, str(value))
    return node


def solid(parent, color, opacity=1):
    node = element('solidFill')
    rgb = element('srgbClr', val=''.join(f'{round(c * 255):02X}' for c in color))
    if opacity < 1:
        rgb.append(element('alpha', val=round(opacity * 100000)))
    node.append(rgb)
    parent.append(node)


def draw_path(slide, drawing, index):
    box = drawing['rect']
    x, y = box.x0, box.y0
    width, height = max(box.width, 0.01), max(box.height, 0.01)
    shape = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, emu(x), emu(y),
                                   emu(width), emu(height))
    shape.name = f'Vector {index:03d}'
    # Remove inherited theme effects (notably the default Office shadow).
    style = shape._element.find('{http://schemas.openxmlformats.org/presentationml/2006/main}style')
    if style is not None:
        shape._element.remove(style)
    props = shape._element.spPr
    for child in list(props):
        if child.tag.rsplit('}', 1)[-1] != 'xfrm':
            props.remove(child)
    geom = element('custGeom')
    for tag in ('avLst', 'gdLst', 'ahLst', 'cxnLst'):
        geom.append(element(tag))
    geom.append(element('rect', l='0', t='0', r='r', b='b'))
    paths = element('pathLst')
    path = element('path', w=emu(width), h=emu(height),
                   fill='norm' if drawing['fill'] is not None else 'none',
                   stroke='1' if drawing['color'] is not None else '0')

    def point_node(point):
        return element('pt', x=emu(point[0] - x), y=emu(point[1] - y))

    def move(point):
        node = element('moveTo')
        node.append(point_node(point))
        path.append(node)

    def line(point):
        node = element('lnTo')
        node.append(point_node(point))
        path.append(node)

    previous = None
    for item in drawing['items']:
        if item[0] == 're':
            r, direction = item[1:]
            points = [r.tl, r.tr, r.br, r.bl] if direction == 1 else [r.tl, r.bl, r.br, r.tr]
            move(points[0])
            for point in points[1:]:
                line(point)
            path.append(element('close'))
            previous = None
        elif item[0] in ('l', 'c'):
            start = item[1]
            if previous is None or abs(previous.x - start.x) + abs(previous.y - start.y) > 0.001:
                move(start)
            if item[0] == 'l':
                line(item[2])
            else:
                node = element('cubicBezTo')
                for point in item[2:]:
                    node.append(point_node(point))
                path.append(node)
            previous = item[-1]
        else:
            raise ValueError(f'Unsupported path item: {item[0]}')
    if drawing['closePath']:
        path.append(element('close'))
    paths.append(path)
    geom.append(paths)
    props.append(geom)
    if drawing['fill'] is None:
        props.append(element('noFill'))
    else:
        solid(props, drawing['fill'], drawing['fill_opacity'])
    stroke = element('ln', w=emu(drawing['width'] or 0), cap='flat')
    if drawing['color'] is None:
        stroke.append(element('noFill'))
    else:
        solid(stroke, drawing['color'], drawing['stroke_opacity'])
        dash = re.search(r'\[\s*([\d.]+)\s+([\d.]+)\s*\]', drawing['dashes'] or '')
        if dash:
            dash_node = element('custDash')
            dash_node.append(element('ds', d=round(float(dash[1]) / drawing['width'] * 100000),
                                     sp=round(float(dash[2]) / drawing['width'] * 100000)))
            stroke.append(dash_node)
        stroke.append(element('round' if drawing['lineJoin'] == 1 else 'miter'))
    props.append(stroke)
    props.append(element('effectLst'))


def draw_text(slide, trace, index):
    chars = trace['chars']
    text = ''.join(chr(c[0]) for c in chars)
    font_name = trace['font']
    # Matplotlib emits this script P at a private-use code point. Use its
    # standard Unicode equivalent to keep the formula editable and portable.
    if '\ue23c' in text:
        text = text.replace('\ue23c', '\U0001D4AB')
        family = 'DejaVu Math TeX Gyre'
    else:
        family = 'DejaVu Sans'
    size = trace['size']
    x, baseline = chars[0][2]
    width = max(trace['bbox'][2] - x, size) + size * 2
    shape = slide.shapes.add_textbox(emu(x), emu(baseline - BASELINE_ASCENT * size),
                                     emu(width), emu(size * 1.6))
    shape.name = f'Text {index:03d}: {text}'
    tf = shape.text_frame
    tf.clear()
    tf.margin_left = tf.margin_right = tf.margin_top = tf.margin_bottom = 0
    tf.word_wrap = False
    tf.vertical_anchor = MSO_ANCHOR.TOP
    tf.auto_size = None
    paragraph = tf.paragraphs[0]
    paragraph.space_before = paragraph.space_after = Pt(0)
    run = paragraph.add_run()
    run.text = text
    run.font.name = family
    run.font.size = Pt(size * SCALE)
    run.font.bold = 'Bold' in font_name
    run.font.italic = 'Oblique' in font_name or ('Italic' in font_name and family == 'DejaVu Sans')
    run.font.color.rgb = RGBColor(*(round(c * 255) for c in trace['color']))
    run._r.get_or_add_rPr().set('kern', '0')


def split_math_spacing(trace):
    """Preserve explicit horizontal gaps emitted by Matplotlib mathtext."""
    groups = [[]]
    for char in trace['chars']:
        if groups[-1] and char[2][0] - groups[-1][-1][3][2] > 0.2:
            groups.append([])
        groups[-1].append(char)
    for group in groups:
        span = dict(trace)
        span['chars'] = group
        span['bbox'] = (group[0][3][0], min(c[3][1] for c in group),
                        group[-1][3][2], max(c[3][3] for c in group))
        yield span


def main():
    doc = pymupdf.open(ROOT / 'spare_architecture.pdf')
    page = doc[0]
    deck = Presentation()
    deck.slide_width = emu(page.rect.width)
    deck.slide_height = emu(page.rect.height)
    deck.core_properties.title = 'SPARE architecture'
    deck.core_properties.subject = 'Editable reconstruction of spare_architecture.pdf'
    slide = deck.slides.add_slide(deck.slide_layouts[6])
    drawings = page.get_drawings()
    texts = [span for trace in page.get_texttrace() for span in split_math_spacing(trace)]
    events = [(d['seqno'], 'path', d) for d in drawings]
    events += [(t['seqno'], 'text', t) for t in texts]
    for index, (_, kind, payload) in enumerate(sorted(events, key=lambda e: e[0])):
        (draw_path if kind == 'path' else draw_text)(slide, payload, index)
    slide.notes_slide.notes_text_frame.text = (
        'Editable vector reconstruction of the selected original architecture PDF. '
        'Text and vector paths are native PowerPoint objects. Math subscripts are '
        'separate text objects. Fonts: DejaVu Sans and DejaVu Math TeX Gyre. '
        'The second slide is the original image for visual comparison.')
    reference = deck.slides.add_slide(deck.slide_layouts[6])
    reference.shapes.add_picture(str(ROOT / 'spare_architecture.png'), 0, 0,
                                 width=deck.slide_width, height=deck.slide_height)
    reference.notes_slide.notes_text_frame.text = 'Original figure, 600 dpi visual reference.'
    output = ROOT / 'spare_architecture.pptx'
    deck.save(output)
    print(json.dumps({'output': str(output), 'slides': 2, 'editable_paths': len(drawings),
                      'editable_text_objects': len(texts)}, indent=2))


if __name__ == '__main__':
    main()
